"""生成 Word 内力表（弯矩 / 剪力）。"""
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from aggregate import StationForces, fmt_val


def _set_cell_shading(cell, hex_color: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _center_cell(cell, text: str, bold: bool = False, size: int = 9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _merge_vertical(table, col: int, row_start: int, row_end: int):
    if row_end <= row_start:
        return
    a = table.cell(row_start, col)
    b = table.cell(row_end, col)
    a.merge(b)


def build_moment_document(
    rows: List[StationForces],
    n_trans: int,
    title: str = "Tableau calculé des efforts internes (Unité : kN × m)",
) -> Document:
    doc = Document()
    doc.add_paragraph(title).alignment = WD_ALIGN_PARAGRAPH.CENTER

    n_cols = 3 + 2 * (1 + n_trans)  # Travée, Pos, Combo + (long + n_trans) × (MAX,MIN)
    header_rows = 2
    n_body = len(rows) * 2
    table = doc.add_table(rows=header_rows + n_body, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # 表头第 1 行
    h0 = [
        "Travée", "Position\nlongitudinale", "Combinaison",
        "Longitudinal\nmoments", "",
    ]
    for i in range(n_trans):
        h0.extend([f"Transversal\nmoments {i + 1}", ""])
    for ci, txt in enumerate(h0):
        _center_cell(table.cell(0, ci), txt, bold=True)
        _set_cell_shading(table.cell(0, ci), "B8CCE4")

    # 表头第 2 行 MAX/MIN
    _center_cell(table.cell(1, 0), "", bold=True)
    _center_cell(table.cell(1, 1), "", bold=True)
    _center_cell(table.cell(1, 2), "", bold=True)
    col = 3
    for _ in range(1 + n_trans):
        _center_cell(table.cell(1, col), "MAX", bold=True)
        _set_cell_shading(table.cell(1, col), "B8CCE4")
        _center_cell(table.cell(1, col + 1), "MIN", bold=True)
        _set_cell_shading(table.cell(1, col + 1), "B8CCE4")
        col += 2

    # 合并表头（先横向再纵向，避免重叠）
    table.cell(0, 3).merge(table.cell(0, 4))
    for i in range(n_trans):
        c0 = 5 + i * 2
        table.cell(0, c0).merge(table.cell(0, c0 + 1))
    for c in range(3):
        table.cell(0, c).merge(table.cell(1, c))

    # 数据
    r = header_rows
    i = 0
    while i < len(rows):
        sf = rows[i]
        travee = sf.station.travee
        # 找同 travée 连续行
        j = i
        while j < len(rows) and rows[j].station.travee == travee:
            j += 1
        travee_row_start = r
        travee_row_end = r + 2 * (j - i) - 1

        for k in range(i, j):
            sfk = rows[k]
            pos_row_start = r
            for combo in ("ELU", "ELS"):
                _center_cell(table.cell(r, 2), combo)
                fields = sfk.combos.get(combo, {})
                long_c = fields.get("long", None)
                if long_c:
                    _center_cell(table.cell(r, 3), fmt_val(long_c.max_val))
                    _center_cell(table.cell(r, 4), fmt_val(long_c.min_val))
                col = 5
                for ti in range(1, n_trans + 1):
                    tc = fields.get(f"trans_{ti}")
                    if tc:
                        _center_cell(table.cell(r, col), fmt_val(tc.max_val))
                        _center_cell(table.cell(r, col + 1), fmt_val(tc.min_val))
                    col += 2
                for c in range(n_cols):
                    _set_cell_shading(table.cell(r, c), "DCE6F1" if combo == "ELU" else "FFFFFF")
                r += 1
            _center_cell(table.cell(pos_row_start, 1), sfk.station.position_label)
            _merge_vertical(table, 1, pos_row_start, r - 1)

        _center_cell(table.cell(travee_row_start, 0), f"Travée {travee}")
        _merge_vertical(table, 0, travee_row_start, travee_row_end)
        i = j

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)

    return doc


def build_shear_document(
    rows: List[StationForces],
    n_trans: int,
    title: str = "Tableau calculé des efforts internes (Unité : kN)",
) -> Document:
    """剪力表：纵桥向 Vxx 合并为一列「Longitudinal cisaillements」，横向各位置单独列。"""
    doc = Document()
    doc.add_paragraph(title).alignment = WD_ALIGN_PARAGRAPH.CENTER

    n_cols = 3 + 2 * (1 + n_trans)
    header_rows = 2
    n_body = len(rows) * 2
    table = doc.add_table(rows=header_rows + n_body, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    h0 = ["Position\nlongitudinale", "", "Combinaison", "Longitudinal\ncisaillements", ""]
    for i in range(n_trans):
        h0.extend([f"Transversal\ncisaillements", ""])
    for ci, txt in enumerate(h0):
        _center_cell(table.cell(0, ci), txt, bold=True)
        _set_cell_shading(table.cell(0, ci), "B8CCE4")

    _center_cell(table.cell(1, 0), "", bold=True)
    _center_cell(table.cell(1, 1), "", bold=True)
    _center_cell(table.cell(1, 2), "", bold=True)
    col = 3
    for _ in range(1 + n_trans):
        _center_cell(table.cell(1, col), "MAX", bold=True)
        _set_cell_shading(table.cell(1, col), "B8CCE4")
        _center_cell(table.cell(1, col + 1), "MIN", bold=True)
        _set_cell_shading(table.cell(1, col + 1), "B8CCE4")
        col += 2

    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))
    table.cell(0, 2).merge(table.cell(1, 2))
    table.cell(0, 3).merge(table.cell(0, 4))
    for i in range(n_trans):
        c0 = 5 + i * 2
        table.cell(0, c0).merge(table.cell(0, c0 + 1))

    r = header_rows
    i = 0
    while i < len(rows):
        sf = rows[i]
        travee = sf.station.travee
        j = i
        while j < len(rows) and rows[j].station.travee == travee:
            j += 1
        block_start = r
        block_end = r + 2 * (j - i) - 1

        for k in range(i, j):
            sfk = rows[k]
            pos_start = r
            for combo in ("ELU", "ELS"):
                _center_cell(table.cell(r, 2), combo)
                fields = sfk.combos.get(combo, {})
                long_c = fields.get("long")
                if long_c:
                    _center_cell(table.cell(r, 3), fmt_val(long_c.max_val))
                    _center_cell(table.cell(r, 4), fmt_val(long_c.min_val))
                col = 5
                for ti in range(1, n_trans + 1):
                    tc = fields.get(f"trans_{ti}")
                    if tc:
                        _center_cell(table.cell(r, col), fmt_val(tc.max_val))
                        _center_cell(table.cell(r, col + 1), fmt_val(tc.min_val))
                    col += 2
                r += 1
            _center_cell(table.cell(pos_start, 1), sfk.station.position_label)
            _merge_vertical(table, 1, pos_start, r - 1)

        _center_cell(table.cell(block_start, 0), f"Travée {travee}")
        _merge_vertical(table, 0, block_start, block_end)
        i = j

    return doc


def save_document(doc: Document, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
