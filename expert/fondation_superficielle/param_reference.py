# -*- coding: utf-8 -*-
"""参数取值表与取值依据 — 供 input/output Excel 计算说明使用"""

from typing import Any, Dict, List, Optional, Tuple

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 与 excel_format 一致的样式（避免循环导入）
FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")
FILL_SECTION = PatternFill("solid", fgColor="E2EFDA")
FILL_TITLE = PatternFill("solid", fgColor="2F5496")
FONT_HEADER = Font(bold=True, size=10)
FONT_SECTION = Font(bold=True, size=11)
FONT_TITLE = Font(bold=True, size=12, color="FFFFFF")
FONT_NORMAL = Font(size=10)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# (参数键, 中文名, 单位, 分类, 取值来源, 规范/公式依据, 备注)
INPUT_PARAM_ROWS: List[Tuple[str, str, str, str, str, str, str]] = [
    ("H_cm", "基础厚度 H", "cm", "几何", "结构总体设计 / 基础施工图", "Fascicule 62 基础几何", "各墩独立填写"),
    ("B_cm", "基础宽度 B", "cm", "几何", "结构总体设计 / 基础施工图", "Fascicule 62；影响 iδβ 中 B/(2De)", "桥台一般大于桥墩"),
    ("L_cm", "基础长度 L", "cm", "几何", "结构总体设计 / 基础施工图", "Fascicule 62 Meyerhof L′=L−2ey", "—"),
    ("De_cm", "基础埋深 De", "cm", "几何", "地面至基底竖向深度", "Fascicule 62 Profondeur d'encastrement；iδβ 用 B/(2De)", "扩大基础，非墩柱直径"),
    ("Zw_cm", "地下水位埋深", "cm", "几何", "水文地质 / 勘察报告", "自地面起算；q₀ 分 γ/γ′", "无地下水填 0"),
    ("gamma_kN_m3", "天然重度 γ", "kN/m³", "土工", "地勘报告", "q₀=γ×D（水位以上）", "—"),
    ("gamma_buoy_kN_m3", "浮重度 γ′", "kN/m³", "土工", "地勘报告", "q₀ 水位以下段", "留空=γ−10"),
    ("q0_kPa", "有效覆土压力 q′₀", "kPa", "土工", "自动：γ×De 或手动", "Fascicule 62 B.3.1.1", "留空自动计算"),
    ("kp", "承载力系数 kp", "—", "土工", "EC7/Fasc.62 特征值表", "q′u=q₀+kp×ple*", "与 B/L、土质有关"),
    ("ple_star_kPa", "极限压力 ple*", "kPa", "土工", "旁压试验区平均或手动", "Annexe E：基底下 0～1.5B 内 pl* 平均", "pl*=plm−σ′v0"),
    ("q_u_kPa", "极限承载力 q′u（可选）", "kPa", "土工", "直接给定则覆盖上式", "q′u=q₀+kp×ple*", "留空则程序自动计算"),
    ("phi_deg", "内摩擦角 φ", "°", "土工", "地勘报告 / 三轴或直剪", "滑动、Kp、被动土压力", "—"),
    ("C_kPa", "有效粘聚力 C′", "kPa", "土工", "地勘报告", "滑动验算限 ≤75 kPa", "默认是否计入见 glissement_include_C"),
    ("gamma_kN_m3", "土体重度 γ", "kN/m³", "土工", "地勘报告", "Fp=½γhp²KpL", "—"),
    ("hp_m", "被动区高度 hp", "m", "滑动", "基础埋深 / 计算书", "Rankine 被动土压力", "留空取 H/100"),
    ("Kp", "被动土压力系数 Kp", "—", "滑动", "计算：tan²(45°+φ/2)", "Fascicule 62 B.3.4 / Word 8.2.2.4", "留空则按 φ 自动算"),
    ("Fp_kN", "被动土压力 Fp", "kN", "滑动", "计算书给定或 ½γhp²KpL", "Word 8.2.2.4", "留空则自动计算"),
    ("glissement_include_C", "滑动是否计 C′项", "0/1", "滑动", "项目约定 / Fascicule 62", "Hd≤Fp+Ntanφ/1.2+C′A′/1.5", "样例计算书多为 0"),
    ("Ec_kPa", "Ménard 模量 Ec", "kPa", "沉降", "旁压 EM 或手动", "Annexe F.2：基底下 0～B/2 层 EM", "Ec=E1"),
    ("Ed_kPa", "Ménard 模量 Ed", "kPa", "沉降", "旁压 EM 五层调和平均", "DTU/Fasc.62 式 2.44", "—"),
    ("alpha_settlement", "流变系数 α", "—", "沉降", "地勘报告 / Fascicule 62 土质表", "Annexe F.2", "砂土常见 0.5"),
    ("B0", "参考宽度 B₀", "m", "沉降", "规范固定值", "Fascicule 62 Annexe F.2", "一般 0.60 m"),
    ("lambda_c", "形状系数 λc", "—", "沉降", "地勘报告 / 基础 L/B", "Annexe F.2，与 L/B 有关", "—"),
    ("lambda_d", "形状系数 λd", "—", "沉降", "地勘报告 / 基础 L/B", "Annexe F.2，与 L/B 有关", "—"),
    ("tau_ulim_kPa", "抗剪极限 τulim", "kPa", "配筋", "混凝土强度等级 / BAEL", "Word 8.2.3 抗剪", "C30/37 等对应限值"),
    ("rho_min_pct", "最小配筋率", "%", "配筋", "规范 / 抗震构造", "Word 8.2.3.1", "样例 0.28%"),
    ("bar_e_cm", "配筋有效高度 e", "cm", "配筋", "配筋构造 HA 间距", "d=H−e（A-A）；计算书 e=15 或 20", "HA25(e=15cm) 等"),
    ("cover_top_cm", "上缘保护层", "cm", "配筋", "配筋构造", "τu=Vd/d，d=(H−保护层)/100", "样例约 9 cm"),
    ("fe_MPa", "钢筋强度 fe", "MPa", "配筋", "材料表 B500B 等", "As=Mu/(0.9·d·fe)", "常用 435 或 500"),
    ("As_adopted_cm2_m", "采用配筋面积", "cm²/m", "配筋", "配筋设计 / 构造", "输出说明用", "HA25@15 等换算"),
]

LOAD_PARAM_ROWS: List[Tuple[str, str, str, str, str, str, str]] = [
    ("depth_m", "旁压试验深度", "m", "旁压", "勘察报告", "自地面起算", "—"),
    ("plm_MPa", "极限压力 plm", "MPa", "旁压", "Ménard 试验", "计算 pl*=plm−σ′v0", "—"),
    ("EM_MPa", "旁压模量 EM", "MPa", "旁压", "Ménard 试验", "计算 Ec/Ed", "—"),
    ("N_kN", "竖向力 N", "kN", "荷载-承载力", "MIDAS Civil 荷载组合提取", "ELU/ELS 各工况", "与 ex、ey 同组合"),
    ("ex_m", "偏心 ex", "m", "荷载-承载力", "MIDAS：M/N 或软件输出", "沿 B 方向；q′ref、iδβ", "—"),
    ("ey_m", "偏心 ey", "m", "荷载-承载力", "MIDAS：M/N 或软件输出", "沿 L 方向；Meyerhof L′", "—"),
    ("Hd_kN", "水平力 Hd", "kN", "荷载-滑动", "MIDAS ELU 组合", "Fascicule 62 B.3.4 滑动", "与 ELU 竖向工况一一对应"),
    ("Mmax_ELU_kNm_m", "弯矩 Mmax ELU", "kN·m/m", "荷载-配筋", "MIDAS 或基础局部计算", "Word 8.2.3 配筋", "截面 A-A / B-B"),
    ("Mmax_ELS_kNm_m", "弯矩 Mmax ELS", "kN·m/m", "荷载-配筋", "MIDAS 或基础局部计算", "准永久 / 稀有", "—"),
    ("Vd_ELU_kN", "剪力 Vd ELU", "kN", "荷载-配筋", "MIDAS 或基础局部计算", "τu=Vd/d", "B-B 可为负值"),
]

FIXED_COEF_ROWS: List[Tuple[str, str, str, str, str]] = [
    ("γq (ELU)", "2.0", "承载力分项", "Fascicule 62 B.2.2.3", "程序内置，不可改"),
    ("γq (ELS)", "3.0", "承载力分项", "Fascicule 62 B.2.2.3", "程序内置"),
    ("γg1", "1.2", "滑动摩擦分项", "Fascicule 62 B.3.4", "N·tanφ 除以 γg1"),
    ("γg2", "1.5", "滑动粘聚力分项", "Fascicule 62 B.3.4", "C′·A′ 除以 γg2"),
    ("Φ1 (iδβ)", "5.3 或 1.0", "偏心折减", "Annexe F.1", "B/(2De)>2.5 →5.3，否则→1.0"),
    ("倾覆 ELU", "As/A≥10%", "倾覆限值", "Fascicule 62 B.3.2", "程序内置"),
    ("倾覆 ELS rare", "As/A≥75%", "倾覆限值", "Fascicule 62 B.3.3", "程序内置"),
    ("ELS QP", "σmin>0", "全截面受压", "Fascicule 62 B.3.3", "qmin>0"),
    ("B0 沉降", "0.6 m", "参考宽度", "Annexe F.2", "除非输入覆盖"),
]

# 输入参数键 → 结果 dict 中实际取值键（若不同）
RESULT_KEY_MAP = {
    "q_u_kPa": "q_u_kPa",
    "phi1_sand": "phi1_sand",
    "B_over_2De": "B_over_2De",
}


def _write_table_header(ws, row: int, headers: List[str], ncol: int) -> int:
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = FILL_HEADER
        c.font = FONT_HEADER
        c.alignment = ALIGN_CENTER
    return row + 1


def _write_section_title(ws, row: int, text: str, ncol: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = FILL_SECTION
    c.font = FONT_SECTION
    c.alignment = ALIGN_LEFT
    return row + 1


def write_param_guide_sheet(ws) -> None:
    """input.xlsx — 参数取值说明（通用参考表）"""
    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = "浅基础验算 — 参数取值表与取值依据"
    c.fill = FILL_TITLE
    c.font = FONT_TITLE
    c.alignment = ALIGN_LEFT

    row = 3
    row = _write_section_title(ws, row, "一、固定系数与验算限值（程序内置，无需输入）", 7)
    row = _write_table_header(ws, row, ["符号/名称", "取值", "用途", "规范依据", "备注", "", ""], 7)
    for name, val, use, basis, note in FIXED_COEF_ROWS:
        for col, v in enumerate([name, val, use, basis, note], 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = FONT_NORMAL
            c.alignment = ALIGN_LEFT
        row += 1

    row += 1
    row = _write_section_title(ws, row, "二、各墩输入参数（在墩柱 sheet 顶部黄色单元格填写）", 7)
    row = _write_table_header(
        ws, row,
        ["参数名", "中文名称", "单位", "分类", "取值来源", "规范/公式依据", "备注"],
        7,
    )
    for key, cn, unit, cat, src, basis, note in INPUT_PARAM_ROWS:
        vals = [key, cn, unit, cat, src, basis, note]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = FONT_NORMAL
            c.alignment = ALIGN_LEFT
        row += 1

    row += 1
    row = _write_section_title(ws, row, "三、荷载与内力（在墩柱 sheet 下方荷载表填写）", 7)
    row = _write_table_header(
        ws, row,
        ["列名", "中文名称", "单位", "分类", "取值来源", "规范/公式依据", "备注"],
        7,
    )
    for key, cn, unit, cat, src, basis, note in LOAD_PARAM_ROWS:
        vals = [key, cn, unit, cat, src, basis, note]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = FONT_NORMAL
            c.alignment = ALIGN_LEFT
        row += 1

    row += 1
    notes = [
        "MIDAS 提取建议：",
        "  · 竖向承载力：各 ELU/ELS 组合下的 N、Mx、My → ex=Mx/N，ey=My/N（注意符号与方向）；",
        "  · 滑动：同序号 ELU 组合的水平力 Hd；",
        "  · 配筋：基础顶面 A-A/B-B 截面 ELU/ELS 弯矩、ELU 剪力（或经局部计算得到）。",
        "桥台与桥墩：B/(2De)>2.5 时 iδβ 用 Φ1=5.3（如墩柱1桥台）；桥墩一般为 Φ1=1.0。",
    ]
    for line in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value=line)
        c.font = FONT_NORMAL
        c.alignment = ALIGN_LEFT
        row += 1

    widths = {1: 18, 2: 16, 3: 8, 4: 10, 5: 28, 6: 32, 7: 24}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _fmt_val(v: Any) -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, float):
        if abs(v - round(v)) < 1e-6:
            return str(int(round(v)))
        return str(round(v, 4))
    return str(v)


def write_output_calc_explanation(wb, results: List[Dict[str, Any]]) -> None:
    """output.xlsx — 计算说明：各墩实际取值 + 取值依据"""
    ws = wb.create_sheet("计算说明", 0)
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "计算说明 — 参数取值表与取值依据"
    c.fill = FILL_TITLE
    c.font = FONT_TITLE
    c.alignment = ALIGN_LEFT

    row = 3
    row = _write_section_title(ws, row, "一、程序内置系数（各墩相同）", 8)
    row = _write_table_header(ws, row, ["名称", "取值", "用途", "规范依据", "备注", "", "", ""], 8)
    for name, val, use, basis, note in FIXED_COEF_ROWS:
        for col, v in enumerate([name, val, use, basis, note], 1):
            ws.cell(row=row, column=col, value=v).alignment = ALIGN_LEFT
        row += 1

    row += 1
    pier_names = [r["pier_name"] for r in results]
    ncol = 4 + len(pier_names)
    row = _write_section_title(ws, row, "二、各墩输入参数实际取值", ncol)
    headers = ["参数名", "中文名", "单位", "取值来源"] + pier_names
    row = _write_table_header(ws, row, headers, ncol)

    key_to_row = {r[0]: r for r in INPUT_PARAM_ROWS}
    for key, cn, unit, _cat, src, _basis, _note in INPUT_PARAM_ROWS:
        vals = [key, cn, unit, src]
        for res in results:
            p = res.get("input_used", {})
            rp = res.get("params", {})
            v = p.get(key)
            if v is None:
                v = rp.get(key.replace("_kPa", "_kPa").replace("_cm", "_cm"))
            if v is None and key in rp:
                v = rp[key]
            if key == "q_u_kPa":
                v = rp.get("q_u_kPa")
            vals.append(_fmt_val(v))
        for col, v in enumerate(vals, 1):
            ws.cell(row=row, column=col, value=v).alignment = ALIGN_LEFT
        row += 1

    # 派生量
    row += 1
    row = _write_section_title(ws, row, "三、程序自动计算 / 派生参数", ncol)
    derived_headers = ["参数名", "中文名", "单位", "说明"] + pier_names
    row = _write_table_header(ws, row, derived_headers, ncol)
    derived_defs = [
        ("q_u_kPa", "极限承载力 q′u", "kPa", "q₀+kp×ple* 或输入覆盖"),
        ("phi1_sand", "iδβ 系数 Φ1", "—", "B/(2De)>2.5→5.3，否则→1.0"),
        ("B_over_2De", "B/(2De)", "—", "判定 Φ1 用"),
        ("Kp_used", "采用 Kp", "—", "输入或 tan²(45°+φ/2)"),
        ("hp_used_m", "采用 hp", "m", "输入或 H"),
        ("q0_calc_kPa", "自动 q′₀=γ×De", "kPa", "γ×De（分 γ/γ′）"),
        ("ple_source", "ple* 来源", "—", "menard / manual"),
        ("Ec_source", "Ec 来源", "—", "menard E1 层"),
        ("Ed_source", "Ed 来源", "—", "menard 五层调和平均"),
    ]
    for key, cn, unit, desc in derived_defs:
        vals = [key, cn, unit, desc]
        for res in results:
            d = res.get("derived", {})
            vals.append(_fmt_val(d.get(key)))
        for col, v in enumerate(vals, 1):
            ws.cell(row=row, column=col, value=v).alignment = ALIGN_LEFT
        row += 1

    row += 1
    row = _write_section_title(ws, row, "四、荷载数据说明", ncol)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row=row, column=1, value=(
        "荷载表（N, ex, ey, Hd, Mmax, Vd）来自 MIDAS 各工况组合，详见 input.xlsx 各墩柱 sheet。"
        " 本 output 各墩 sheet 中表格为程序按 Fascicule 62 公式计算结果。"
    ))
    c.alignment = ALIGN_LEFT
    row += 2

    row = _write_section_title(ws, row, "五、完整参数取值依据（参考）", 8)
    row = _write_table_header(
        ws, row,
        ["参数名", "中文名", "单位", "分类", "取值来源", "规范/公式依据", "备注", ""],
        8,
    )
    for key, cn, unit, cat, src, basis, note in INPUT_PARAM_ROWS + LOAD_PARAM_ROWS:
        vals = [key, cn, unit, cat, src, basis, note, ""]
        for col, v in enumerate(vals, 1):
            ws.cell(row=row, column=col, value=v).alignment = ALIGN_LEFT
        row += 1

    widths = {1: 16, 2: 14, 3: 8, 4: 22, 5: 28, 6: 30, 7: 20, 8: 14}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
